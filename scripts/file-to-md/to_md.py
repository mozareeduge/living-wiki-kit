"""to_md.py - convert content files to clean markdown with provenance header.

Usage:
    python to_md.py "C:/path/file.pdf" [-o "C:/path/file.extracted.md"]

Prints a JSON summary to stdout. Exit 0 on success, 2 on unsupported type.
Requires: pymupdf (pdf/epub), python-docx (docx), python-pptx (pptx),
          openpyxl (xlsx), beautifulsoup4+lxml (html). No network, no models.
"""
import argparse
import datetime
import json
import pathlib
import sys

MAX_XLSX_ROWS = 200
MIN_PAGE_CHARS = 50
FA_CHARS = set("ابپتثجچحخدرزژسشصضطظعغفقکگلمنوهیيىةك")


def fa_ratio(text):
    if not text.strip():
        return 0.0
    return sum(1 for c in text if c in FA_CHARS) / max(len(text), 1)


def quality_note(text):
    if fa_ratio(text) > 0.05:
        return "fa-spacing: Persian glyphs may extract letter-spaced; verify quotes vs source"
    return "clean"


def header(src, method, detail=""):
    line = ("# Extracted from " + src.name + "\n\n> method: " + method + " | bytes: "
            + str(src.stat().st_size) + " | date: " + datetime.date.today().isoformat())
    if detail:
        line += " | " + detail
    return line + "\n\n"


def md_table(rows):
    norm = [[str(c).replace("|", "\\|").replace(chr(10), " ").strip() for c in r] for r in rows]
    width = max(len(r) for r in norm)
    norm = [r + [""] * (width - len(r)) for r in norm]
    out = ["| " + " | ".join(norm[0]) + " |",
           "| " + " | ".join(["---"] * width) + " |"]
    out += ["| " + " | ".join(r) + " |" for r in norm[1:]]
    return "\n".join(out)


def from_pdf(src):
    import pymupdf
    doc = pymupdf.open(str(src))
    parts, scanned = [], []
    for i, page in enumerate(doc):
        t = page.get_text().strip()
        if len(t) < MIN_PAGE_CHARS:
            scanned.append(i + 1)
            continue
        parts.append("## Page " + str(i + 1) + "\n\n" + t)
    body = "\n\n".join(parts)
    detail = "pages: " + str(len(doc))
    if scanned:
        detail += " | no-text-layer pages (needs OCR): " + str(scanned)
    needs_ocr = len(body.strip()) < 100
    return body, "pymupdf-text", detail, needs_ocr


def from_docx(src):
    from docx import Document
    d = Document(str(src))
    parts = []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        st = p.style.name if p.style else ""
        if st.startswith("Heading 1"):
            parts.append("# " + t)
        elif st.startswith("Heading 2"):
            parts.append("## " + t)
        elif st.startswith("Heading 3"):
            parts.append("### " + t)
        elif "List" in st:
            parts.append("- " + t)
        else:
            parts.append(t)
    for n, t in enumerate(d.tables):
        rows = [[c.text.strip() for c in row.cells] for row in t.rows]
        rows = [r for r in rows if any(r)]
        if rows:
            parts.append("### Table " + str(n + 1) + "\n\n" + md_table(rows))
    detail = "paragraphs: " + str(len(parts)) + " | tables: " + str(len(d.tables))
    return "\n\n".join(parts), "python-docx", detail, False


def from_pptx(src):
    from pptx import Presentation
    prs = Presentation(str(src))
    parts = []
    for i, slide in enumerate(prs.slides):
        title = ""
        try:
            if slide.shapes.title and slide.shapes.title.text.strip():
                title = slide.shapes.title.text.strip()
        except Exception:
            title = ""
        head = "## Slide " + str(i + 1) + ((": " + title) if title else "")
        parts.append(head)
        for shape in slide.shapes:
            if shape.has_table:
                rows = [[c.text.strip() for c in row.cells] for row in shape.table.rows]
                rows = [r for r in rows if any(r)]
                if rows:
                    parts.append(md_table(rows))
            elif shape.has_text_frame and shape.text.strip():
                if shape.text.strip() != title:
                    parts.append(shape.text.strip())
        try:
            notes = slide.notes_slide.placeholders[1].text.strip()
            if notes:
                parts.append("*Notes: " + notes + "*")
        except Exception:
            pass
    return "\n\n".join(parts), "python-pptx", "slides: " + str(len(prs.slides)), False


def from_xlsx(src):
    import openpyxl
    wb = openpyxl.load_workbook(str(src), data_only=True, read_only=True)
    parts, truncated = [], []
    for ws in wb.worksheets:
        rows = [[("" if c is None else str(c)) for c in row]
                for row in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(x.strip() for x in r)]
        if not rows:
            continue
        if len(rows) > MAX_XLSX_ROWS:
            truncated.append(ws.title + " (" + str(len(rows)) + " rows, kept " + str(MAX_XLSX_ROWS) + ")")
            rows = rows[:MAX_XLSX_ROWS]
        parts.append("## Sheet: " + ws.title + "\n\n" + md_table(rows))
    detail = "sheets: " + str(len(wb.worksheets))
    if truncated:
        detail += " | truncated: " + "; ".join(truncated)
    return "\n\n".join(parts), "openpyxl", detail, False


def from_html(src):
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(src.read_bytes(), "lxml")
    for tag in soup(["script", "style", "nav", "footer"]):
        tag.decompose()
    parts = []
    body = soup.body or soup
    for el in body.find_all(["h1", "h2", "h3", "p", "li", "table"]):
        if el.name == "h1":
            parts.append("# " + el.get_text(" ", strip=True))
        elif el.name == "h2":
            parts.append("## " + el.get_text(" ", strip=True))
        elif el.name == "h3":
            parts.append("### " + el.get_text(" ", strip=True))
        elif el.name == "li":
            parts.append("- " + el.get_text(" ", strip=True))
        elif el.name == "table":
            rows = [[c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"])]
                    for tr in el.find_all("tr")]
            rows = [r for r in rows if any(r)]
            if rows:
                parts.append(md_table(rows))
        else:
            t = el.get_text(" ", strip=True)
            if t:
                parts.append(t)
    if not parts:
        t = soup.get_text(chr(10), strip=True)
        parts = [t] if t else []
    return "\n\n".join(parts), "beautifulsoup", "", False


def from_text(src):
    return src.read_text(encoding="utf-8", errors="replace").strip(), "passthrough", "", False


HANDLERS = {
    ".pdf": from_pdf, ".epub": from_pdf,
    ".docx": from_docx,
    ".pptx": from_pptx,
    ".xlsx": from_xlsx, ".csv": from_text,
    ".html": from_html, ".htm": from_html,
    ".txt": from_text, ".md": from_text, ".markdown": from_text,
}
REFUSED = {
    ".doc": "legacy .doc - export to .docx and retry",
    ".xls": "legacy .xls - export to .xlsx and retry",
    ".ppt": "legacy .ppt - export to .pptx and retry",
    ".odt": "no LibreOffice here - export to .docx and retry",
    ".rtf": "no converter installed - export to .docx/.txt and retry",
    ".zip": "archive - unpack first, then convert members",
}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("source")
    ap.add_argument("-o", "--output", default=None)
    args = ap.parse_args()
    src = pathlib.Path(args.source)
    if not src.is_file():
        print(json.dumps({"ok": False, "error": "not found: " + args.source}))
        return 1
    ext = src.suffix.lower()
    if ext in REFUSED:
        print(json.dumps({"ok": False, "error": REFUSED[ext]}))
        return 2
    if ext not in HANDLERS:
        print(json.dumps({"ok": False, "error": "unsupported type: " + ext}))
        return 2
    try:
        body, method, detail, needs_ocr = HANDLERS[ext](src)
    except Exception as e:
        print(json.dumps({"ok": False, "error": "extract failed (" + ext + "): " + str(e)}))
        return 1
    note = quality_note(body)
    if needs_ocr:
        extra = "" if note == "clean" else "; " + note
        note = "NEEDS-OCR: no text layer; do not quote this file; route per skill" + extra
    md = header(src, method, detail)
    if note != "clean" or needs_ocr:
        md += "> quality: " + note + "\n\n"
    md += body + "\n"
    out = pathlib.Path(args.output) if args.output else src.parent / (src.stem + ".extracted.md")
    out.write_text(md, encoding="utf-8")
    print(json.dumps({"ok": True, "output": str(out), "chars": len(md),
                      "tokens_est": len(md) // 4, "method": method,
                      "needs_ocr": needs_ocr, "quality": note}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
